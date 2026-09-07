import os
import fitz  # PyMuPDF
import docx

class ExtractionError(Exception):
    pass

def extract_text(file_path):
    """
    Extracts text and metadata from the given file based on its extension.
    
    :param file_path: Absolute path to the file
    :return: A dictionary containing:
             - 'text': Pure text content (string)
             - 'lines': Text split by line (list of strings)
             - 'formatting': Dictionary containing details of fonts/styles if available
    """
    if not os.path.exists(file_path):
        raise ExtractionError("File does not exist.")
        
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.docx':
        return _extract_docx(file_path)
    elif ext == '.pdf':
        return _extract_pdf(file_path)
    elif ext in ['.jpg', '.jpeg', '.png']:
        return {
            'text': '',
            'lines': [],
            'formatting': {
                'image_detected': True,
                'fonts_detected': ['Scanned Layout']
            }
        }
    else:
        raise ExtractionError(f"Unsupported file format: {ext}")

def _extract_docx(file_path):
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        raise ExtractionError(f"Failed to open DOCX file: {str(e)}")
        
    lines = []
    formatting_info = {
        'paragraphs_count': len(doc.paragraphs),
        'tables_count': len(doc.tables),
        'styles_detected': []
    }
    
    styles = set()
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
            # Gather basic styles from paragraph
            if para.style and para.style.name:
                styles.add(para.style.name)
                
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                lines.append(" | ".join(row_cells))
                
    full_text = "\n".join(lines)
    if not full_text.strip():
        raise ExtractionError("The DOCX document is empty.")
        
    formatting_info['styles_detected'] = list(styles)
    
    return {
        'text': full_text,
        'lines': lines,
        'formatting': formatting_info
    }

def _extract_pdf(file_path):
    try:
        doc = fitz.open(file_path)
    except Exception as e:
        raise ExtractionError(f"Failed to open PDF file: {str(e)}")
        
    lines = []
    formatting_info = {
        'pages_count': len(doc),
        'fonts_detected': []
    }
    
    fonts = set()
    for page in doc:
        # Extract plain text
        text = page.get_text()
        if text.strip():
            page_lines = [line.strip() for line in text.split('\n') if line.strip()]
            lines.extend(page_lines)
            
        # Extract formatting fonts if available
        try:
            page_dict = page.get_text("dict")
            for block in page_dict.get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        font_name = span.get("font", "")
                        if font_name:
                            fonts.add(font_name)
        except Exception:
            # Fallback if dictionary extraction fails on complex layouts
            pass
            
    full_text = "\n".join(lines)
    # Note: If no text is found, we do not throw. We allow empty text returns so the system
    # can trigger multimodal OCR processing via Gemini (Phase 3).
    if not full_text.strip():
        formatting_info['fonts_detected'] = ['Scanned Layout / Image-only PDF']
    else:
        formatting_info['fonts_detected'] = list(fonts)
    
    return {
        'text': full_text,
        'lines': lines,
        'formatting': formatting_info
    }
