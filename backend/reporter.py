import os
import json
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ReportLab imports for PDF
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'uploads')

def _set_cell_background(cell, fill_hex):
    """Helper to set cell background color in docx."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def generate_docx_report(record):
    """
    Generates a Word Document (.docx) report for a comparison record.
    
    :param record: Comparison record dict containing file names, date, summary, and results.
    :return: Absolute path to the generated DOCX file
    """
    doc_id = record['id']
    file_a = record['file_a_name']
    file_b = record['file_b_name']
    summary = record['summary']
    changes = record['results'].get('changes', [])
    
    doc = Document()
    
    # Style definitions
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(11)
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Contract Comparison Audit Report")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    
    # Subtitle/Date
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(f"Generated on {record.get('comparison_date', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # 1. Executive Summary Heading
    h1 = doc.add_paragraph()
    h1_run = h1.add_run("1. Executive Verdict")
    h1_run.font.size = Pt(16)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(15, 23, 42)
    
    # Verdict Panel
    verdict_box = doc.add_table(rows=1, cols=1)
    verdict_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = verdict_box.cell(0, 0)
    _set_cell_background(cell, "F1F5F9") # light gray shading
    
    v_para = cell.paragraphs[0]
    v_para.paragraph_format.left_indent = Inches(0.1)
    v_para.paragraph_format.right_indent = Inches(0.1)
    v_run = v_para.add_run(summary.get('verdict', 'No verdict generated.'))
    v_run.font.italic = True
    v_run.font.size = Pt(11)
    
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.add_run(f"Original Document: ").bold = True
    p.add_run(f"{file_a}\n")
    p.add_run(f"Modified Document: ").bold = True
    p.add_run(f"{file_b}\n")
    
    # 2. Key Metrics Heading
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2_run = h2.add_run("2. Audit Statistics")
    h2_run.font.size = Pt(16)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(15, 23, 42)
    
    # Metrics Table
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Shading Accent 1'
    
    metrics = [
        ("Textual Modifications", summary.get('textualChanges', 0)),
        ("Grammatical Corrections", summary.get('grammarChanges', 0)),
        ("Formatting Updates", summary.get('formattingChanges', 0)),
        ("Date Discrepancies", summary.get('dateChanges', 0)),
        ("Visual Elements Checked", summary.get('visualChanges', 0))
    ]
    
    for i, (name, val) in enumerate(metrics):
        table.rows[i].cells[0].text = name
        table.rows[i].cells[1].text = str(val)
        table.rows[i].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # 3. Detailed Changes List
    h3 = doc.add_paragraph()
    h3_run = h3.add_run("3. Detailed List of Detected Changes")
    h3_run.font.size = Pt(16)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(15, 23, 42)
    
    if not changes:
        doc.add_paragraph("No specific changes were logged in the detailed report.")
    else:
        changes_table = doc.add_table(rows=1, cols=5)
        changes_table.style = 'Table Grid'
        
        # Headers
        hdr_cells = changes_table.rows[0].cells
        hdr_cells[0].text = 'Category'
        hdr_cells[1].text = 'Severity'
        hdr_cells[2].text = 'Section'
        hdr_cells[3].text = 'Description'
        hdr_cells[4].text = 'Original ➔ Modified Text'
        
        for cell in hdr_cells:
            _set_cell_background(cell, "4F46E5") # Indigo header
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            cell.paragraphs[0].runs[0].font.bold = True
            
        for c in changes:
            row_cells = changes_table.add_row().cells
            row_cells[0].text = c.get('category', 'Textual')
            row_cells[1].text = c.get('severity', 'Low')
            row_cells[2].text = c.get('section', 'General')
            row_cells[3].text = c.get('description', '')
            
            orig = c.get('originalText', '')
            mod = c.get('modifiedText', '')
            
            p_diff = row_cells[4].paragraphs[0]
            if orig:
                r_orig = p_diff.add_run(f"[-] {orig}\n")
                r_orig.font.color.rgb = RGBColor(239, 68, 68)
            if mod:
                r_mod = p_diff.add_run(f"[+] {mod}")
                r_mod.font.color.rgb = RGBColor(16, 185, 129)
                
            # Formatting severities visually
            sev = c.get('severity', 'Low')
            if sev == 'High':
                _set_cell_background(row_cells[1], "FEE2E2") # red background for High
                row_cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(185, 28, 28)
                row_cells[1].paragraphs[0].runs[0].font.bold = True
            elif sev == 'Medium':
                _set_cell_background(row_cells[1], "FEF3C7") # yellow background for Medium
                row_cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 83, 9)
                row_cells[1].paragraphs[0].runs[0].font.bold = True

    # Save DOCX
    docx_filename = f"comparison_report_{doc_id}.docx"
    docx_path = os.path.join(UPLOAD_FOLDER, docx_filename)
    doc.save(docx_path)
    
    return docx_path

def generate_pdf_report(record):
    """
    Generates a PDF report using ReportLab.
    
    :param record: Comparison record dict
    :return: Absolute path to the generated PDF file
    """
    doc_id = record['id']
    file_a = record['file_a_name']
    file_b = record['file_b_name']
    summary = record['summary']
    changes = record['results'].get('changes', [])
    
    pdf_filename = f"comparison_report_{doc_id}.pdf"
    pdf_path = os.path.join(UPLOAD_FOLDER, pdf_filename)
    
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        leftMargin=0.5*inch,
        rightMargin=0.5*inch,
        topMargin=0.5*inch,
        bottomMargin=0.5*inch
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#4F46E5'),
        alignment=1, # Center
        spaceAfter=6
    )
    
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        textColor=colors.HexColor('#64748B'),
        alignment=1, # Center
        spaceAfter=18
    )
    
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=14,
        spaceAfter=8
    )
    
    body_style = ParagraphStyle(
        'ReportBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155')
    )
    
    verdict_style = ParagraphStyle(
        'VerdictBody',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=11,
        leading=16,
        textColor=colors.HexColor('#1E293B')
    )
    
    diff_style = ParagraphStyle(
        'DiffText',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=10
    )

    story = []
    
    # Header Banner
    story.append(Paragraph("Contract Comparison Audit Report", title_style))
    story.append(Paragraph(f"Generated on {record.get('comparison_date', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}", meta_style))
    
    # 1. Verdict Section
    story.append(Paragraph("1. Executive Verdict", heading_style))
    
    verdict_text = summary.get('verdict', 'No verdict generated.')
    verdict_table = Table([[Paragraph(verdict_text, verdict_style)]], colWidths=[7.0*inch])
    verdict_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F1F5F9')),
        ('TOPPADDING', (0,0), (-1,-1), 12),
        ('BOTTOMPADDING', (0,0), (-1,-1), 12),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(verdict_table)
    story.append(Spacer(1, 12))
    
    # Files details text
    files_text = f"<b>Original Contract:</b> {file_a}<br/><b>Modified Contract:</b> {file_b}"
    story.append(Paragraph(files_text, body_style))
    story.append(Spacer(1, 12))
    
    # 2. Stats Section
    story.append(Paragraph("2. Audit Statistics", heading_style))
    
    stats_data = [
        ["Category", "Changes Count"],
        ["Textual Modifications", str(summary.get('textualChanges', 0))],
        ["Grammatical Corrections", str(summary.get('grammarChanges', 0))],
        ["Formatting Updates", str(summary.get('formattingChanges', 0))],
        ["Date Discrepancies", str(summary.get('dateChanges', 0))],
        ["Visual Elements Checked", str(summary.get('visualChanges', 0))]
    ]
    
    stats_table = Table(stats_data, colWidths=[4.0*inch, 3.0*inch])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F8FAFC')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#475569')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 8),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('ALIGN', (1,0), (1,-1), 'RIGHT'),
        ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(stats_table)
    story.append(Spacer(1, 16))
    
    # 3. Changes Section
    story.append(Paragraph("3. Detailed Comparison Log", heading_style))
    
    if not changes:
        story.append(Paragraph("No detailed differences were recorded in this log.", body_style))
    else:
        # Table Headers
        changes_data = [[
            Paragraph("<b>Category</b>", body_style),
            Paragraph("<b>Severity</b>", body_style),
            Paragraph("<b>Section</b>", body_style),
            Paragraph("<b>Description</b>", body_style),
            Paragraph("<b>Original ➔ Modified</b>", body_style)
        ]]
        
        for c in changes:
            # Build Category Cell
            cat = c.get('category', 'Textual')
            
            # Build Severity Cell
            sev = c.get('severity', 'Low')
            sev_color = '#F3F4F6' # Low (light gray)
            if sev == 'High':
                sev_color = '#FEE2E2' # Red
            elif sev == 'Medium':
                sev_color = '#FEF3C7' # Yellow
            
            # Build Diff Cell
            orig = c.get('originalText', '')
            mod = c.get('modifiedText', '')
            
            diff_paragraphs = []
            if orig:
                diff_paragraphs.append(f"<font color='#EF4444'>[-] {orig}</font>")
            if mod:
                diff_paragraphs.append(f"<font color='#10B981'>[+] {mod}</font>")
            
            diff_text = "<br/>".join(diff_paragraphs)
            
            changes_data.append([
                Paragraph(cat, body_style),
                Paragraph(f"<font color='#1F2937'><b>{sev}</b></font>", body_style),
                Paragraph(c.get('section', 'General'), body_style),
                Paragraph(c.get('description', ''), body_style),
                Paragraph(diff_text, diff_style)
            ])
            
        changes_table = Table(changes_data, colWidths=[1.1*inch, 0.8*inch, 1.3*inch, 2.0*inch, 2.3*inch])
        
        # Build style list
        t_style = [
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#4F46E5')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ]
        
        # Color headers text white
        for col_idx in range(5):
            t_style.append(('TEXTCOLOR', (col_idx, 0), (col_idx, 0), colors.white))
            
        # Color cells individually for High/Medium severities
        for row_idx, c in enumerate(changes, start=1):
            sev = c.get('severity', 'Low')
            if sev == 'High':
                t_style.append(('BACKGROUND', (1, row_idx), (1, row_idx), colors.HexColor('#FEE2E2')))
            elif sev == 'Medium':
                t_style.append(('BACKGROUND', (1, row_idx), (1, row_idx), colors.HexColor('#FEF3C7')))
                
        changes_table.setStyle(TableStyle(t_style))
        story.append(changes_table)
        
    doc.build(story)
    return pdf_path
